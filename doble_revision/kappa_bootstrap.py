"""
IC 95 % del kappa de Cohen por bootstrap percentil (remuestreo de eventos con reemplazo).
Lee las respuestas ORIGINALES (antes del consenso) del maestro y escribe el resultado en la hoja «IC bootstrap».

  pip install openpyxl numpy
  python kappa_bootstrap.py maestro_analisis_kappa.xlsx
"""
import sys, datetime
import numpy as np
from openpyxl import load_workbook

L = "ABCDEFG"; B = 10_000; SEMILLA = 2026; MIN_VALIDAS = 0.95  # mínimo de réplicas con kappa definido


def leer(wb):
    col = lambda h: [str(c.value).strip().upper() if c.value not in (None, "") else "" for c in wb[h]["B"][1:29]]
    r1, r2 = col("R1 respuestas"), col("R2 respuestas")
    malos = [(i + 1, a or "(vacío)", b or "(vacío)") for i, (a, b) in enumerate(zip(r1, r2)) if a not in L or b not in L]
    if malos:
        sys.exit("No se calculó nada. Corrige estos eventos (evento, R1, R2): " + str(malos))
    return np.array([L.index(x) for x in r1]), np.array([L.index(x) for x in r2])


def kappa(a, b):
    po = np.mean(a == b)
    pe = sum(np.mean(a == k) * np.mean(b == k) for k in range(7))
    return np.nan if np.isclose(pe, 1.0) else (po - pe) / (1 - pe)


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else "maestro_analisis_kappa.xlsx"
    wb = load_workbook(ruta); a, b = leer(wb)
    n = len(a); po = float(np.mean(a == b)); k = kappa(a, b)
    res = {"n eventos": n, "Acuerdo observado": round(po, 4)}
    if not np.isfinite(k):
        res.update({"Kappa": "indefinido (pe = 1)", "IC 95 % inferior": "—", "IC 95 % superior": "—", "Réplicas válidas": "—",
                    "Método": "Kappa indefinido: todas las respuestas en una sola causa.",
                    "Frase para el paper": f"Ambos revisores asignaron la misma única causa a los {n} eventos (acuerdo observado {po:.0%}); el kappa de Cohen no está definido."})
    else:
        rng = np.random.default_rng(SEMILLA)
        reps = np.array([kappa(a[i], b[i]) for i in (rng.integers(0, n, n) for _ in range(B))])
        val = reps[np.isfinite(reps)]
        res.update({"Kappa": round(float(k), 4), "Réplicas válidas": f"{len(val)} de {B} (semilla {SEMILLA})"})
        if len(val) < MIN_VALIDAS * B:
            res.update({"IC 95 % inferior": "—", "IC 95 % superior": "—",
                        "Método": f"Bootstrap no fiable: menos del {MIN_VALIDAS:.0%} de réplicas con kappa definido.",
                        "Frase para el paper": f"κ = {k:.2f} (n = {n}; acuerdo observado {po:.0%}); IC no estimable por bootstrap con esta distribución de causas."})
        else:
            lo, hi = np.percentile(val, [2.5, 97.5])
            res.update({"IC 95 % inferior": round(float(lo), 4), "IC 95 % superior": round(float(hi), 4),
                        "Método": f"Bootstrap percentil, remuestreo de eventos con reemplazo, {B} réplicas",
                        "Frase para el paper": f"κ = {k:.2f} (IC 95 %: {lo:.2f}–{hi:.2f}; bootstrap percentil, {B} réplicas; n = {n}; acuerdo observado {po:.0%})."})
    res["Fecha de cálculo"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    hoja = wb["IC bootstrap"]
    for fila in range(3, 12):
        clave = hoja.cell(fila, 1).value
        if clave in res: hoja.cell(fila, 2, res[clave])
    wb.save(ruta)
    for c, v in res.items(): print(f"{c}: {v}")
    print(f"\nGuardado en la hoja «IC bootstrap» de {ruta}")


if __name__ == "__main__":
    main()
